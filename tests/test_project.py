import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

import app
import backend
import audit_controller as mcp
import rag_pipeline as rag
from models import Base, engine as application_engine
from reports import ReportGenerator


class FakeEmbeddings(Embeddings):
    def embed_documents(self, texts):
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        return [float("approval" in text.lower()), float("risk" in text.lower()), 1.0]


def tearDownModule():
    application_engine.dispose()


class FakeRetriever:
    def invoke(self, question):
        return [Document(page_content="Manager approval is required.", metadata={"source": "policy.txt"})]


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        engine = create_engine("sqlite://")
        self.addCleanup(engine.dispose)
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine, expire_on_commit=False)
        for target in ("backend.SessionLocal", "audit_controller.SessionLocal"):
            patcher = patch(target, factory)
            patcher.start()
            self.addCleanup(patcher.stop)
        config = patch.dict(app.app.config, TESTING=True, UPLOAD_FOLDER=str(self.folder))
        config.start()
        self.addCleanup(config.stop)
        self.client = app.app.test_client()

    def test_dashboard_and_form_routes(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"FRMC Copilot Dashboard", page.data)
        with self.client.get("/static/style.css") as response:
            self.assertEqual(response.status_code, 200)
        with patch("app.get_retriever", return_value=FakeRetriever()):
            self.assertEqual(self.client.post("/ask", data={"question": "Approval?"}).status_code, 200)
            result = self.client.post("/generate_report", data={"question": "Approval?"})
            self.assertEqual(result.status_code, 201)
        self.assertEqual(len(self.client.get("/reports").json), 1)
        self.assertEqual(len(self.client.get("/logs").json), 1)

    def test_builtin_answers_and_agents_without_provider(self):
        with patch("rag_pipeline.get_embeddings", side_effect=AssertionError("Unexpected API usage")):
            response = self.client.post("/query", json={"question": "What are the five COSO components?"})
            self.assertEqual(response.status_code, 200)
            self.assertIn("Monitoring Activities", response.json["answer"])
            response = self.client.post("/run_agents", json={
                "question": "What is SOX?",
                "evidence": {"amount": 600000, "approval": False},
                "risk_factors": {"criticality": "High"},
            })
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json["risk_result"], "NOT_ASSESSED")
            self.assertEqual(response.json["agent_trace"], ["RAG", "Policy & Control", "Evidence", "Risk", "Review"])
            self.assertEqual(len(self.client.get("/logs").json), 1)

    def test_invalid_requests(self):
        for route in ("/query", "/run_agents", "/report", "/generate_report"):
            for body in (None, [], {}, {"question": 12}, {"question": " "}):
                with self.subTest(route=route, body=body):
                    result = self.client.post(route, data=json.dumps(body), content_type="application/json")
                    self.assertEqual(result.status_code, 400)
        self.assertEqual(self.client.post("/query", data="{", content_type="application/json").status_code, 400)
        self.assertEqual(self.client.get("/missing").status_code, 404)

    def test_report_transactions_and_compatibility(self):
        self.assertEqual(self.client.post("/report", json={"title": "T"}).status_code, 400)
        result = self.client.post("/report", json={"title": "T", "content": "C"})
        self.assertEqual(result.status_code, 201)
        report = ReportGenerator().export_report(result.json["id"])
        self.assertEqual(report.content, "C")
        with self.assertRaises(Exception):
            with backend.SessionLocal.begin() as session:
                from models import Report
                session.add(Report(title=None, content="bad"))
        self.assertEqual(backend.report_service.create_report("Next", "Valid").title, "Next")

    def test_workflow_serialization_and_inputs(self):
        with patch("app.get_retriever", return_value=FakeRetriever()):
            result = self.client.post("/run_agents", json={
                "question": "Approval?", "evidence": {"amount": 600000, "approval": False},
                "risk_factors": {"criticality": "High", "exception": True},
            })
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json["final_report"]["risk"], "NOT_ASSESSED")
        self.assertTrue(result.json["final_report"]["evidence"])
        self.assertEqual(result.json["final_report"]["evidence_assessment"]["status"], "NOT_ASSESSED")
        self.assertEqual(result.json["rag_result"], [])
        self.assertNotIn("retriever", result.json)
        for values in ({"evidence": None}, {"evidence": {"amount": "large"}},
                       {"evidence": {"approval": "false"}}, {"risk_factors": []},
                       {"risk_factors": {"exception": "false"}}):
            self.assertEqual(self.client.post("/run_agents", json={"question": "Q", **values}).status_code, 400)

    def test_upload_validation_and_cleanup(self):
        self.assertEqual(self.client.post("/upload").status_code, 400)
        for name, contents in (("bad.exe", b"x"), ("empty.txt", b""), ("bad.pdf", b"not a pdf")):
            result = self.client.post("/upload", data={"file": (io.BytesIO(contents), name)})
            self.assertEqual(result.status_code, 400)
        self.assertEqual(list(self.folder.iterdir()), [])
        with patch("app.build_vectorstore", side_effect=rag.KnowledgeBaseError("Not configured")):
            self.assertEqual(self.client.post("/upload", data={"file": (io.BytesIO(b"Financial audit policy"), "a.txt")}).status_code, 409)
        self.assertFalse((self.folder / "a.txt").exists())

    def test_upload_sanitizes_and_prevents_overwrite(self):
        with patch("app.build_vectorstore") as build:
            result = self.client.post("/upload", data={"file": (io.BytesIO(b"Financial approval policy"), "../../policy.TXT")})
            self.assertEqual(result.status_code, 200)
            self.assertEqual(build.call_args.args[0][0].metadata["source"], "policy.TXT")
            result = self.client.post("/upload", data={"file": (io.BytesIO(b"Changed"), "policy.TXT")})
            self.assertEqual(result.status_code, 409)
        self.assertEqual((self.folder / "policy.TXT").read_text(), "Financial approval policy")

    def test_empty_knowledge_base(self):
        with patch("rag_pipeline.CHROMA_DIR", str(self.folder / "missing")):
            result = self.client.post("/query", json={"question": "unrelated astronomy"} )
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.json["answer_mode"], "insufficient_knowledge")
        self.assertEqual(self.client.post("/build_rag").status_code, 400)

    def test_provider_failure_response(self):
        from openai import APIConnectionError
        import httpx
        failure = APIConnectionError(request=httpx.Request("POST", "https://example.invalid"))
        with patch("app.UploadedKnowledge", side_effect=failure):
            self.assertEqual(self.client.post("/query", json={"question": "Q"}).status_code, 502)


class ExtractionTests(unittest.TestCase):
    def test_text_csv_excel_docx(self):
        import pandas as pd
        import docx
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "data.TXT").write_text("Approval policy", encoding="utf-8")
            (root / "data.csv").write_text("amount,approval\n600000,false", encoding="utf-8")
            with pd.ExcelWriter(root / "data.xlsx") as writer:
                pd.DataFrame({"first": [1]}).to_excel(writer, sheet_name="First")
                pd.DataFrame({"second": [2]}).to_excel(writer, sheet_name="Second")
            document = docx.Document()
            document.add_paragraph("Policy")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
            document.save(root / "data.docx")
            (root / "subfolder").mkdir()
            self.assertEqual(len(rag.load_documents(root, validate_finance=False)), 4)
            self.assertIn("Second", rag.extract_text_from_file(root / "data.xlsx"))
            self.assertIn("Table evidence", rag.extract_text_from_file(root / "data.docx"))
            self.assertEqual(rag.load_documents(root / "missing"), [])

    def test_embeddings_work_without_api_key(self):
        rag.get_embeddings.cache_clear()
        vector = rag.get_embeddings().embed_query("Financial audit controls")
        self.assertEqual(len(vector), 1024)
        self.assertTrue(any(vector))


class VectorTests(unittest.TestCase):
    def test_rebuild_retrieval_and_stale_chunk_removal(self):
        # Chroma retains file handles on Windows; explicitly close its test system.
        import chromadb
        from chromadb.config import Settings
        from langchain_chroma import Chroma
        with tempfile.TemporaryDirectory() as folder:
            client = chromadb.PersistentClient(path=folder, settings=Settings(anonymized_telemetry=False))
            try:
                store = Chroma(client=client, embedding_function=FakeEmbeddings())
                docs = [Document(page_content="Approval policy", metadata={"source": "a.txt"})]
                with patch("rag_pipeline.open_vectorstore", return_value=store):
                    rag.build_vectorstore(docs)
                    rag.build_vectorstore(docs)
                    self.assertEqual(len(store.get()["ids"]), 1)
                    self.assertEqual(store.as_retriever().invoke("approval")[0].page_content, "Approval policy")
                    rag.build_vectorstore([Document(page_content="Risk policy", metadata={"source": "a.txt"})])
                    self.assertEqual(store.get()["documents"], ["Risk policy"])
                    rag.build_vectorstore([
                        Document(page_content="Risk policy", metadata={"source": "a.txt"}),
                        Document(page_content="Financial audit", metadata={"source": "b.txt"}),
                    ])
                    with patch("rag_pipeline.CHROMA_DIR", folder):
                        rag.remove_document_from_index("a.txt")
                        self.assertEqual(store.get()["documents"], ["Financial audit"])
                        rag.remove_document_from_index("b.txt")
                        self.assertEqual(store.get()["ids"], [])
                        self.assertEqual(rag.get_retriever().invoke("unrelated astronomy"), [])
            finally:
                client._system.stop()


if __name__ == "__main__":
    unittest.main()
