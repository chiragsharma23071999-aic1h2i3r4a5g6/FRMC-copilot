import hashlib
from functools import lru_cache
from pathlib import Path

import pandas as pd
import docx
import PyPDF2
from local_embeddings import LocalLexicalEmbeddings
from langchain_chroma import Chroma
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from config import CHROMA_DIR, UPLOAD_DIR
from knowledge import reference_documents
from document_policy import require_finance_content

SUPPORTED_EXTENSIONS = {".txt", ".csv", ".xlsx", ".xls", ".pdf", ".docx"}


class KnowledgeBaseError(RuntimeError):
    pass


@lru_cache(maxsize=1)
def get_embeddings():
    return LocalLexicalEmbeddings()


def extract_text_from_file(filepath):
    path = Path(filepath)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file format: {suffix}")
    try:
        if suffix == ".txt":
            return path.read_text(encoding="utf-8-sig")
        if suffix == ".csv":
            return pd.read_csv(path).to_string(index=False)
        if suffix in {".xlsx", ".xls"}:
            sheets = pd.read_excel(path, sheet_name=None, engine="xlrd" if suffix == ".xls" else "openpyxl")
            return "\n\n".join(f"{name}\n{frame.to_string(index=False)}" for name, frame in sheets.items())
        if suffix == ".pdf":
            with path.open("rb") as stream:
                return "\n".join(page.extract_text() or "" for page in PyPDF2.PdfReader(stream).pages)
        document = docx.Document(path)
        paragraphs = [p.text for p in document.paragraphs]
        paragraphs.extend(" | ".join(cell.text for cell in row.cells)
                          for table in document.tables for row in table.rows)
        return "\n".join(paragraphs)
    except Exception as exc:
        raise ValueError(f"Could not read {path.name}. Check the file format and contents.") from exc


def load_documents(upload_folder=UPLOAD_DIR, validate_finance=True):
    folder = Path(upload_folder)
    if not folder.exists():
        return []
    docs = []
    for path in sorted(folder.iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            text = extract_text_from_file(path)
            if validate_finance:
                require_finance_content(text, path.name)
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": path.name}))
    return docs


def open_vectorstore():
    return Chroma(collection_name="frmc_local_lexical_v1", persist_directory=CHROMA_DIR,
                  embedding_function=get_embeddings(), collection_metadata={"hnsw:space": "cosine"})


def build_vectorstore(docs):
    chunks = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50).split_documents(docs)
    if not chunks:
        raise ValueError("No readable text found. Upload a document containing text first.")
    store = open_vectorstore()
    # Stable IDs prevent duplicates when the same documents are rebuilt.
    ids = [hashlib.sha256(f"{doc.metadata.get('source', '')}\0{i}\0{doc.page_content}".encode()).hexdigest()
           for i, doc in enumerate(chunks)]
    old_ids = set(store.get()["ids"])
    for start in range(0, len(chunks), 256):
        store.add_documents(chunks[start:start + 256], ids=ids[start:start + 256])
    stale_ids = old_ids - set(ids)
    if stale_ids:
        store.delete(ids=list(stale_ids))
    return store


def remove_document_from_index(filename):
    if Path(CHROMA_DIR).exists():
        open_vectorstore().delete(where={"source": filename})


def get_upload_retriever():
    if not Path(CHROMA_DIR).exists():
        raise KnowledgeBaseError("Upload and process a document before querying the knowledge base.")
    store = open_vectorstore()
    if not store.get(limit=1)["ids"]:
        raise KnowledgeBaseError("Upload and process a document before querying the knowledge base.")
    return store.as_retriever(search_kwargs={"k": 3})


class ProjectRetriever:
    def invoke(self, question):
        references = reference_documents(question)
        if references:
            return references
        if not Path(CHROMA_DIR).exists():
            return []
        try:
            return get_upload_retriever().invoke(question)
        except KnowledgeBaseError:
            return []


def get_retriever():
    return ProjectRetriever()
